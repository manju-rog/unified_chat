#!/usr/bin/env python3
"""
Test SOW generation with the exact warehouse inventory input
"""

import requests
import json
import time

NEW_SOW_URL = "http://localhost:8002"

def test_warehouse_sow():
    """Test with exact warehouse inventory input"""
    
    print("=" * 80)
    print("TESTING WAREHOUSE INVENTORY SOW GENERATION")
    print("=" * 80)
    
    # Exact input from user
    test_data = {
        "template_path": "/Users/manju/Desktop/project 3/sample_sow_template.docx",
        "project_data": {
            "project_info": "SOW-2025-091, Warehouse Inventory Management System, objectives are automate stock tracking, reduce manual data entry errors, enable real-time inventory visibility, integrate with existing ERP system",
            
            "services": "standard",
            
            "deliverables": "Requirements document with business workflows and user stories. System design specifications with database schema and API definitions. Web dashboard for inventory management and reporting. Mobile scanner app for iOS and Android. ERP integration module with automated sync. Testing documentation and user training materials.",
            
            "timeline": "Start date 2026-02-01, end date 2026-06-30, 10 sprints of 2 weeks each",
            
            "resources": "2 Full-Stack Developers from Engineering team full-time, 1 Mobile Developer from Mobile team full-time, 1 Database Administrator from Data team part-time at 50%, 1 QA Engineer from Quality team full-time, 1 Project Manager from PMO full-time",
            
            "contacts": "Contractor: James Wilson, LogiTech Solutions, 450 Commerce Street Denver CO 80202, +1-303-555-0156, james.wilson@logitech-sol.com, Operations Manager. Client: Maria Garcia, WareMax Distribution, 890 Industrial Parkway Chicago IL 60601, +1-312-555-0423, maria.garcia@waremax.com, Supply Chain Director",
            
            "budget": "Requirements and Design milestone 32000 dollars at sprint 3, Web Application milestone 68000 dollars at sprint 6, Mobile App milestone 45000 dollars at sprint 8, ERP Integration milestone 38000 dollars at sprint 9, Testing and Training milestone 27000 dollars at sprint 10, estimated expenses 12000 dollars for cloud hosting and licenses"
        },
        "session_id": f"warehouse_test_{int(time.time())}"
    }
    
    print("\n📤 Sending request to new_sow...")
    print(f"   Project: Warehouse Inventory Management System")
    print(f"   Services: standard")
    print(f"   Timeline: 10 sprints")
    print(f"   Resources: 6 team members")
    
    try:
        response = requests.post(
            f"{NEW_SOW_URL}/api/generate-direct",
            json=test_data,
            timeout=120
        )
        
        if response.status_code == 200:
            result = response.json()
            print(f"\n✅ Document generated successfully!")
            print(f"   📄 Filename: {result.get('filename')}")
            
            # Check the document
            import os
            from docx import Document
            
            doc_path = f"/Users/manju/Desktop/project 3/new_sow/output/{result.get('filename')}"
            if os.path.exists(doc_path):
                doc = Document(doc_path)
                full_text = "\n".join([para.text for para in doc.paragraphs])
                
                print(f"\n📋 VERIFICATION:")
                
                # Check deliverables
                if "Requirements document" in full_text:
                    print(f"   ✅ Deliverables found")
                    
                    # Check for sprint allocation
                    if "Sprints 1-2" in full_text or "Sprint 1" in full_text:
                        print(f"   ✅ Sprint allocation present")
                    else:
                        print(f"   ⚠️  Sprint allocation missing or showing 'None'")
                else:
                    print(f"   ❌ Deliverables not found")
                
                # Check resources
                if "Full-Stack Developer" in full_text or "Engineering" in full_text:
                    print(f"   ✅ Resources found")
                else:
                    print(f"   ❌ Resources not found")
                
                # Check standard services
                if "Design Review" in full_text and "File Transfer mechanism" in full_text:
                    print(f"   ✅ Standard services applied")
                else:
                    print(f"   ❌ Standard services missing")
                
                # Find and display deliverables table
                print(f"\n📊 DELIVERABLES TABLE:")
                for table in doc.tables:
                    # Check if this is the deliverables table
                    if len(table.rows) > 0:
                        first_row = [cell.text.strip() for cell in table.rows[0].cells]
                        if "Deliverable Name" in " ".join(first_row) or "Deliverable" in " ".join(first_row):
                            print(f"   Found deliverables table with {len(table.rows)} rows")
                            for i, row in enumerate(table.rows[:4]):  # Show first 4 rows
                                row_text = " | ".join([cell.text.strip()[:50] for cell in row.cells])
                                print(f"   Row {i}: {row_text}")
                            break
                
                # Find and display resources
                print(f"\n👥 RESOURCES:")
                in_resources = False
                for para in doc.paragraphs:
                    text = para.text.strip()
                    if "resource" in text.lower() and len(text) < 50:
                        in_resources = True
                    elif in_resources and text:
                        if "Full-Stack Developer" in text or "Mobile Developer" in text or "QA Engineer" in text:
                            print(f"   {text[:100]}")
                        if len(text) > 100 or "milestone" in text.lower():
                            break
                
                return True
            else:
                print(f"   ❌ Document not found: {doc_path}")
                return False
        else:
            print(f"\n❌ Generation failed: {response.status_code}")
            print(f"   Error: {response.text}")
            return False
            
    except Exception as e:
        print(f"\n❌ Error: {e}")
        import traceback
        traceback.print_exc()
        return False

if __name__ == "__main__":
    success = test_warehouse_sow()
    
    print("\n" + "=" * 80)
    if success:
        print("✅ TEST COMPLETED - Check output above for details")
    else:
        print("❌ TEST FAILED")
    print("=" * 80)
