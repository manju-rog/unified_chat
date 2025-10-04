"""
SOW (Statement of Work) Tool Endpoints.
Provides REST APIs for creating, updating, and generating SOW documents.
Requirements: 10.1, 10.2, 10.3
"""

from fastapi import APIRouter, HTTPException
from typing import Dict
from datetime import datetime
import uuid

from app.models import SowStart, SowUpdate, ContactInfo, ServiceItem, DeliverableItem

# ============================================================================
# Router Setup (Subtask 5.1)
# ============================================================================

router = APIRouter(prefix="/sow", tags=["sow"])

# In-memory storage for SOW sessions
# Key: sow_id -> Value: {sow_id, project_name, oracle_rep, billing_contact, services, deliverables, acceptance, created_at, updated_at}
sow_storage: Dict[str, Dict] = {}


# ============================================================================
# Helper Functions
# ============================================================================

def _create_sow_session(sow_id: str, project_name: str) -> Dict:
    """Create a new SOW session in storage"""
    session = {
        "sow_id": sow_id,
        "project_name": project_name,
        "oracle_rep": None,
        "billing_contact": None,
        "services": [],
        "deliverables": [],
        "acceptance": None,
        "created_at": datetime.now().isoformat(),
        "updated_at": datetime.now().isoformat()
    }
    sow_storage[sow_id] = session
    return session


# ============================================================================
# Endpoints (Subtasks 5.2, 5.3, 5.4)
# ============================================================================

@router.post("/start")
async def start_sow(request: SowStart) -> Dict:
    """
    Start a new SOW session.
    
    Requirements: 4.1, 10.1
    
    Args:
        request: SowStart with project_name
        
    Returns:
        {ok: bool, sow_id: str, message: str}
    """
    try:
        # Generate unique SOW ID
        sow_id = str(uuid.uuid4())
        
        # Create SOW session in memory
        _create_sow_session(sow_id, request.project_name)
        
        message = f"Started SOW session for project '{request.project_name}'. SOW ID: {sow_id}"
        
        return {
            "ok": True,
            "sow_id": sow_id,
            "message": message
        }
    
    except Exception as e:
        return {
            "ok": False,
            "reason": f"Failed to start SOW session: {str(e)}"
        }



@router.post("/update")
async def update_sow(request: SowUpdate) -> Dict:
    """
    Update SOW session with collected information.
    
    Requirements: 4.2, 10.2
    
    Args:
        request: SowUpdate with sow_id and optional fields (oracle_rep, billing_contact, services, deliverables, acceptance)
        
    Returns:
        {ok: bool, message: str}
    """
    try:
        # Check if SOW session exists
        if request.sow_id not in sow_storage:
            return {
                "ok": False,
                "reason": f"SOW session {request.sow_id} not found"
            }
        
        session = sow_storage[request.sow_id]
        
        # Update fields if provided
        if request.oracle_rep is not None:
            session["oracle_rep"] = request.oracle_rep.model_dump()
        
        if request.billing_contact is not None:
            session["billing_contact"] = request.billing_contact.model_dump()
        
        if request.services is not None:
            session["services"] = [s.model_dump() for s in request.services]
        
        if request.deliverables is not None:
            session["deliverables"] = [d.model_dump() for d in request.deliverables]
        
        if request.acceptance is not None:
            session["acceptance"] = request.acceptance
        
        # Update timestamp
        session["updated_at"] = datetime.now().isoformat()
        
        message = f"Updated SOW session {request.sow_id}"
        
        return {
            "ok": True,
            "message": message
        }
    
    except Exception as e:
        return {
            "ok": False,
            "reason": f"Failed to update SOW session: {str(e)}"
        }



@router.post("/generate/{sow_id}")
async def generate_sow(sow_id: str) -> Dict:
    """
    Generate a DOCX file from SOW session data.
    
    Requirements: 4.5, 10.3
    
    Args:
        sow_id: SOW session identifier
        
    Returns:
        {ok: bool, url: str}
    """
    try:
        # Check if SOW session exists
        if sow_id not in sow_storage:
            return {
                "ok": False,
                "reason": f"SOW session {sow_id} not found"
            }
        
        session = sow_storage[sow_id]
        
        # Import python-docx
        try:
            from docx import Document
            from docx.shared import Pt, Inches
        except ImportError:
            return {
                "ok": False,
                "reason": "python-docx library not installed. Run: pip install python-docx"
            }
        
        # Create DOCX document
        doc = Document()
        
        # Add title
        title = doc.add_heading('Statement of Work', 0)
        title.alignment = 1  # Center alignment
        
        # Add project name
        doc.add_heading(f"Project: {session['project_name']}", level=1)
        
        # Add Oracle Representative section
        if session.get('oracle_rep'):
            doc.add_heading('Oracle Representative', level=2)
            oracle_rep = session['oracle_rep']
            doc.add_paragraph(f"Name: {oracle_rep.get('name', 'N/A')}")
            doc.add_paragraph(f"Email: {oracle_rep.get('email', 'N/A')}")
            if oracle_rep.get('phone'):
                doc.add_paragraph(f"Phone: {oracle_rep['phone']}")
            if oracle_rep.get('address'):
                doc.add_paragraph(f"Address: {oracle_rep['address']}")
        
        # Add Billing Contact section
        if session.get('billing_contact'):
            doc.add_heading('Billing Contact', level=2)
            billing = session['billing_contact']
            doc.add_paragraph(f"Name: {billing.get('name', 'N/A')}")
            doc.add_paragraph(f"Email: {billing.get('email', 'N/A')}")
            if billing.get('phone'):
                doc.add_paragraph(f"Phone: {billing['phone']}")
            if billing.get('address'):
                doc.add_paragraph(f"Address: {billing['address']}")
        
        # Add Services section
        if session.get('services') and len(session['services']) > 0:
            doc.add_heading('Services', level=2)
            for idx, service in enumerate(session['services'], 1):
                doc.add_paragraph(f"{idx}. {service.get('name', 'N/A')}", style='List Number')
                doc.add_paragraph(f"   {service.get('description', 'N/A')}")
        
        # Add Deliverables section
        if session.get('deliverables') and len(session['deliverables']) > 0:
            doc.add_heading('Deliverables', level=2)
            for idx, deliverable in enumerate(session['deliverables'], 1):
                doc.add_paragraph(f"{idx}. {deliverable.get('name', 'N/A')}", style='List Number')
                doc.add_paragraph(f"   {deliverable.get('description', 'N/A')}")
        
        # Add Acceptance Criteria section
        if session.get('acceptance'):
            doc.add_heading('Acceptance Criteria', level=2)
            doc.add_paragraph(session['acceptance'])
        
        # Add footer with generation timestamp
        doc.add_paragraph()
        footer = doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        footer.runs[0].italic = True
        
        # Save to /tmp directory
        import os
        output_dir = "/tmp"
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
        
        file_path = os.path.join(output_dir, f"{sow_id}.docx")
        doc.save(file_path)
        
        # Return download URL
        url = f"/files/{sow_id}.docx"
        
        return {
            "ok": True,
            "url": url,
            "message": f"SOW document generated successfully"
        }
    
    except Exception as e:
        return {
            "ok": False,
            "reason": f"Failed to generate SOW document: {str(e)}"
        }
