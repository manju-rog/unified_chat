import logging
import os
import shutil
from pathlib import Path
from typing import Optional
from datetime import datetime
from docx import Document
from ..models.sow_models import SOWContext

logger = logging.getLogger(__name__)

class DocumentService:
    """Service for generating SOW documents"""
    
    def __init__(self):
        self.output_dir = Path("output")
        self.output_dir.mkdir(exist_ok=True)
    
    def generate_document(
        self,
        template_path: str,
        sow_context: SOWContext,
        output_dir: Optional[str] = None
    ) -> str:
        """Generate SOW document from template and context"""
        try:
            # Load template
            doc = Document(template_path)
            
            # Replace placeholders
            self._replace_placeholders(doc, sow_context)
            
            # Generate output filename
            project_name = sow_context.project_info.project_name if sow_context.project_info else "SOW"
            safe_name = "".join(c for c in project_name if c.isalnum() or c in (' ', '-', '_')).rstrip()
            safe_name = safe_name.replace(' ', '_')
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{safe_name}_{timestamp}.docx"
            
            # Set output path
            if output_dir:
                output_path = Path(output_dir) / filename
            else:
                output_path = self.output_dir / filename
            
            # Save document
            doc.save(str(output_path))
            
            logger.info(f"Document generated: {output_path}")
            return str(output_path)
            
        except Exception as e:
            logger.error(f"Error generating document: {e}")
            raise
    
    def _replace_placeholders(self, doc: Document, context: SOWContext):
        """Replace placeholders in document with actual data"""
        try:
            # Create replacement dictionary
            replacements = self._build_replacements(context)
            
            # Replace in paragraphs
            for paragraph in doc.paragraphs:
                for placeholder, value in replacements.items():
                    if placeholder in paragraph.text:
                        paragraph.text = paragraph.text.replace(placeholder, str(value))
            
            # Replace in tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for placeholder, value in replacements.items():
                            if placeholder in cell.text:
                                cell.text = cell.text.replace(placeholder, str(value))
            
        except Exception as e:
            logger.error(f"Error replacing placeholders: {e}")
            raise
    
    def _build_replacements(self, context: SOWContext) -> dict:
        """Build dictionary of placeholder replacements"""
        replacements = {}
        
        # Project info
        if context.project_info:
            replacements.update({
                "{{DOCUMENT_NUMBER}}": context.project_info.document_number or "TBD",
                "{{PROJECT_NAME}}": context.project_info.project_name or "TBD",
                "{{OBJECTIVES}}": "\n".join(f"• {obj}" for obj in context.project_info.objectives) if context.project_info.objectives else "TBD"
            })
        
        # Services
        if context.services:
            services_text = ""
            for i, service in enumerate(context.services, 1):
                services_text += f"{i}. {service.name}\n"
                if service.description:
                    services_text += f"   {service.description}\n"
                if service.duration:
                    services_text += f"   Duration: {service.duration}\n"
                services_text += "\n"
            replacements["{{SERVICES}}"] = services_text.strip()
        
        # Deliverables
        if context.deliverables:
            deliverables_text = ""
            for deliv in context.deliverables:
                deliverables_text += f"{deliv.id}. {deliv.name}\n"
                deliverables_text += f"   {deliv.description}\n\n"
            replacements["{{DELIVERABLES}}"] = deliverables_text.strip()
        
        # Timeline
        if context.timeline:
            replacements.update({
                "{{START_DATE}}": context.timeline.start_date.strftime("%Y-%m-%d") if context.timeline.start_date else "TBD",
                "{{END_DATE}}": context.timeline.end_date.strftime("%Y-%m-%d") if context.timeline.end_date else "TBD",
                "{{TOTAL_SPRINTS}}": str(context.timeline.total_sprints) if context.timeline.total_sprints else "TBD",
                "{{SPRINT_DURATION}}": context.timeline.sprint_duration or "2 weeks"
            })
        
        # Resources
        if context.resources:
            resources_text = ""
            for resource in context.resources:
                resources_text += f"• {resource.role} ({resource.team}): {resource.count} person(s), {resource.allocation}\n"
            replacements["{{RESOURCES}}"] = resources_text.strip()
        
        # Contacts
        if context.contractor_contact:
            replacements.update({
                "{{CONTRACTOR_NAME}}": context.contractor_contact.name or "TBD",
                "{{CONTRACTOR_COMPANY}}": context.contractor_contact.company or "TBD",
                "{{CONTRACTOR_ADDRESS}}": context.contractor_contact.address or "TBD",
                "{{CONTRACTOR_PHONE}}": context.contractor_contact.phone or "TBD",
                "{{CONTRACTOR_EMAIL}}": context.contractor_contact.email or "TBD",
                "{{CONTRACTOR_ROLE}}": context.contractor_contact.role or "TBD"
            })
        
        if context.client_contact:
            replacements.update({
                "{{CLIENT_NAME}}": context.client_contact.name or "TBD",
                "{{CLIENT_COMPANY}}": context.client_contact.company or "TBD",
                "{{CLIENT_ADDRESS}}": context.client_contact.address or "TBD",
                "{{CLIENT_PHONE}}": context.client_contact.phone or "TBD",
                "{{CLIENT_EMAIL}}": context.client_contact.email or "TBD",
                "{{CLIENT_ROLE}}": context.client_contact.role or "TBD"
            })
        
        # Budget
        if context.milestones:
            milestones_text = ""
            for milestone in context.milestones:
                milestones_text += f"{milestone.id}. {milestone.name}: ${milestone.fee:,.2f}\n"
                if milestone.description:
                    milestones_text += f"   {milestone.description}\n"
            replacements["{{MILESTONES}}"] = milestones_text.strip()
        
        replacements.update({
            "{{TOTAL_FEE}}": f"${context.total_fee:,.2f}",
            "{{ESTIMATED_EXPENSES}}": f"${context.estimated_expenses:,.2f}"
        })
        
        # Default values for missing placeholders
        default_replacements = {
            "{{DOCUMENT_NUMBER}}": "TBD",
            "{{PROJECT_NAME}}": "TBD",
            "{{OBJECTIVES}}": "TBD",
            "{{SERVICES}}": "TBD",
            "{{DELIVERABLES}}": "TBD",
            "{{START_DATE}}": "TBD",
            "{{END_DATE}}": "TBD",
            "{{TOTAL_SPRINTS}}": "TBD",
            "{{SPRINT_DURATION}}": "2 weeks",
            "{{RESOURCES}}": "TBD",
            "{{CONTRACTOR_NAME}}": "TBD",
            "{{CONTRACTOR_COMPANY}}": "TBD",
            "{{CONTRACTOR_ADDRESS}}": "TBD",
            "{{CONTRACTOR_PHONE}}": "TBD",
            "{{CONTRACTOR_EMAIL}}": "TBD",
            "{{CONTRACTOR_ROLE}}": "TBD",
            "{{CLIENT_NAME}}": "TBD",
            "{{CLIENT_COMPANY}}": "TBD",
            "{{CLIENT_ADDRESS}}": "TBD",
            "{{CLIENT_PHONE}}": "TBD",
            "{{CLIENT_EMAIL}}": "TBD",
            "{{CLIENT_ROLE}}": "TBD",
            "{{MILESTONES}}": "TBD",
            "{{TOTAL_FEE}}": "$0.00",
            "{{ESTIMATED_EXPENSES}}": "$0.00"
        }
        
        # Add defaults for any missing keys
        for key, value in default_replacements.items():
            if key not in replacements:
                replacements[key] = value
        
        return replacements