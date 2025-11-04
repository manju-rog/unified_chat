from pathlib import Path
from typing import Optional, Dict, Any
from datetime import datetime
from docx import Document

try:
    from docxtpl import DocxTemplate
    HAS_DOXCTPL = True
except Exception:
    HAS_DOXCTPL = False

class DocumentService:
    def __init__(self, out_dir: Path):
        self.out_dir = Path(out_dir)
        self.out_dir.mkdir(parents=True, exist_ok=True)
    
    def _default_filename(self, project_name: str = "SOW") -> str:
        """Generate filename with project name and timestamp"""
        safe_name = project_name.replace(" ", "_").replace("/", "_")[:50]  # Limit length
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return f"SOW_{safe_name}_doc_{timestamp}.docx"
    
    def _prepare_template_context(self, context: Dict[str, Any]) -> Dict[str, Any]:
        """Prepare context for template rendering - matching new_sow format"""
        from datetime import datetime
        
        # Extract project info
        project_info = context.get('project_info', 'Professional Services Project')
        
        # Parse services
        services_text = context.get('services', '')
        services_list = []
        if services_text.startswith("SERVICES:"):
            # Parse standard package format
            services_raw = services_text.replace("SERVICES:", "").strip()
            for service in services_raw.split(","):
                service = service.strip()
                if "(" in service and ")" in service:
                    name = service.split("(")[0].strip()
                    duration = service.split("(")[1].split(")")[0].strip()
                    services_list.append({
                        'name': name,
                        'description': f"{name} - {duration}",
                        'duration': duration
                    })
        else:
            # Custom services
            services_list.append({
                'name': 'Custom Services',
                'description': services_text,
                'duration': 'As per project requirements'
            })
        
        # Parse deliverables
        deliverables_text = context.get('deliverables', '')
        deliverables_list = []
        for idx, deliv in enumerate(deliverables_text.split('\n'), 1):
            deliv = deliv.strip().lstrip('•-*').strip()
            if deliv:
                deliverables_list.append({
                    'id': idx,
                    'name': deliv,
                    'description': deliv,
                    'sprint_start': 1,
                    'sprint_end': 2,
                    'sprint_duration': '2 weeks'
                })
        
        # Parse timeline
        timeline_text = context.get('timeline', '')
        
        # Parse resources
        resources = context.get('resources', [])
        resources_list = []
        for r in resources:
            resources_list.append({
                'role': r.get('role', 'Team Member'),
                'count': r.get('count', 1),
                'rate': 'As per agreement'
            })
        
        # Parse contacts
        contacts = context.get('contacts', {})
        
        # Prepare comprehensive context matching new_sow template structure
        template_context = {
            # Document metadata
            "document_number": f"SOW-{datetime.now().strftime('%Y-%m-%d-%H%M')}",
            "current_date": datetime.now().strftime("%B %d, %Y"),
            
            # Project information
            "project_name": project_info,
            "objectives": [project_info],  # Can be enhanced
            
            # Services
            "services": services_list,
            
            # Deliverables
            "deliverables": deliverables_list,
            
            # Timeline
            "start_date": datetime.now().strftime("%B %d, %Y"),
            "end_date": (datetime.now().replace(month=datetime.now().month + 3) if datetime.now().month <= 9 
                        else datetime.now().replace(year=datetime.now().year + 1, month=datetime.now().month - 9)).strftime("%B %d, %Y"),
            "total_sprints": len(deliverables_list) if deliverables_list else 4,
            "sprint_duration": "2 weeks",
            
            # Resources
            "resources": resources_list,
            
            # Client contact
            "client_name": contacts.get('name', 'Client Organization'),
            "client_company": contacts.get('name', 'Client Organization'),
            "client_address": contacts.get('address', 'Client Address'),
            "client_phone": contacts.get('phone', 'Phone Number'),
            "client_email": contacts.get('email', 'client@example.com'),
            "client_role": contacts.get('designation', 'Contact Person'),
            
            # Contractor contact (default values)
            "contractor_name": "Professional Services Team",
            "contractor_company": "Professional Services Team",
            "contractor_address": "123 Business Ave, Suite 100, City, State 12345",
            "contractor_phone": "+1 (555) 123-4567",
            "contractor_email": "pm@company.com",
            "contractor_role": "Project Manager",
            
            # Milestones and budget
            "milestones": [
                {"id": 1, "name": "Project Kickoff", "description": "Initial project setup", "payment_percentage": 25},
                {"id": 2, "name": "Mid-project Review", "description": "Progress review", "payment_percentage": 50},
                {"id": 3, "name": "Final Delivery", "description": "Project completion", "payment_percentage": 100}
            ],
            "total_fee": context.get('budget', 'To be determined'),
            "estimated_expenses": "$0.00",
            
            # Additional
            "assumptions": [
                "Client will provide timely feedback and approvals",
                "All necessary resources and access will be provided",
                "Project scope remains as defined in this SOW"
            ],
            "terms": [
                "Changes to scope require written approval",
                "Payment terms are Net 30 days",
                "Intellectual property rights as per master agreement"
            ]
        }
        
        return template_context
    
    def generate(self,
                context: Dict[str, Any],
                session_id: str,
                template_path: Optional[Path] = None) -> Path:
        
        print(f"📄 Generating SOW document...")
        print(f"📁 Output directory: {self.out_dir}")
        print(f"📋 Template: {template_path}")
        
        # Get project name for filename
        project_name = context.get('project_info', 'SOW')
        if isinstance(project_name, str) and len(project_name) > 50:
            project_name = project_name[:50]
        
        # Prefer a docxtpl template (mirrors new_sow behavior)
        if template_path and template_path.exists() and HAS_DOXCTPL:
            try:
                print(f"✅ Using professional template: {template_path.name}")
                doc = DocxTemplate(str(template_path))
                
                # Prepare context in new_sow format
                template_context = self._prepare_template_context(context)
                
                # Render template
                doc.render(template_context)
                
                # Save to generated_docs_sow folder (not in session subfolder)
                filename = self._default_filename(project_name)
                out = self.out_dir / filename
                doc.save(str(out))
                
                print(f"✅ Professional SOW document generated: {out}")
                print(f"📄 Filename: {filename}")
                return out
            except Exception as e:
                print(f"❌ Template rendering failed: {e}")
                import traceback
                traceback.print_exc()
                print("🔄 Falling back to basic document generation...")
        
        # Enhanced fallback: professional python-docx build with AI content
        print("📝 Creating professional SOW document with enhanced formatting...")
        doc = Document()
        
        # Title page
        title = doc.add_heading('STATEMENT OF WORK', 0)
        title.alignment = 1  # Center alignment
        
        doc.add_paragraph(f"Document Number: {context.get('document_number', 'SOW-001')}")
        doc.add_paragraph(f"Date: {context.get('current_date', 'Current Date')}")
        doc.add_paragraph("")  # Blank line
        
        # Executive Summary
        doc.add_heading('EXECUTIVE SUMMARY', 1)
        doc.add_paragraph(context.get("executive_summary", "Professional services engagement as outlined below."))
        
        # Project Information
        doc.add_heading('1. PROJECT SCOPE & OBJECTIVES', 1)
        doc.add_paragraph(context.get("project_scope", context.get("project_info", "Project scope to be defined.")))
        
        # Services
        doc.add_heading('2. SERVICES DESCRIPTION', 1)
        doc.add_paragraph(context.get("services_description", f"Professional {context.get('services', 'standard')} services."))
        
        # Deliverables
        doc.add_heading('3. DELIVERABLES', 1)
        doc.add_paragraph(context.get("deliverables_description", context.get("deliverables", "Deliverables to be specified.")))
        
        # Timeline
        doc.add_heading('4. TIMELINE & MILESTONES', 1)
        doc.add_paragraph(context.get("timeline_description", context.get("timeline", "Timeline to be established.")))
        
        # Resources
        doc.add_heading('5. RESOURCE ALLOCATION', 1)
        doc.add_paragraph(context.get("team_description", "Professional team to be assigned."))
        if context.get("resources"):
            doc.add_paragraph("Team Composition:")
            for r in context["resources"]:
                doc.add_paragraph(f"• {r['role']}: {r['count']} person(s)", style='List Bullet')
        
        # Client Information
        doc.add_heading('6. CLIENT INFORMATION', 1)
        doc.add_paragraph(f"Organization: {context.get('client_name', 'Client Organization')}")
        doc.add_paragraph(f"Contact Person: {context.get('client_contact_person', 'Contact Person')}")
        doc.add_paragraph(f"Email: {context.get('client_email', 'client@example.com')}")
        doc.add_paragraph(f"Phone: {context.get('client_phone', 'Phone Number')}")
        doc.add_paragraph(f"Address: {context.get('client_address', 'Client Address')}")
        
        # Contractor Information
        doc.add_heading('7. CONTRACTOR INFORMATION', 1)
        doc.add_paragraph(f"Organization: {context.get('contractor_name', 'Professional Services Team')}")
        doc.add_paragraph(f"Contact Person: {context.get('contractor_contact_person', 'Project Manager')}")
        doc.add_paragraph(f"Email: {context.get('contractor_email', 'pm@company.com')}")
        doc.add_paragraph(f"Phone: {context.get('contractor_phone', '+1 (555) 123-4567')}")
        
        # Budget
        doc.add_heading('8. BUDGET & PAYMENT TERMS', 1)
        doc.add_paragraph(f"Total Project Fee: {context.get('total_fee', context.get('budget', 'To be determined'))}")
        doc.add_paragraph(f"Payment Terms: {context.get('payment_terms', 'Net 30 days')}")
        doc.add_paragraph(context.get("budget_description", "Budget breakdown available upon request."))
        
        # Assumptions
        doc.add_heading('9. ASSUMPTIONS', 1)
        assumptions = context.get("assumptions", [
            "Client will provide timely feedback and approvals",
            "All necessary resources and access will be provided",
            "Project scope remains as defined in this SOW"
        ])
        for assumption in assumptions:
            doc.add_paragraph(f"• {assumption}", style='List Bullet')
        
        # Terms
        doc.add_heading('10. TERMS & CONDITIONS', 1)
        terms = context.get("terms", [
            "Changes to scope require written approval",
            "Payment terms are Net 30 days",
            "Intellectual property rights as per master agreement"
        ])
        for term in terms:
            doc.add_paragraph(f"• {term}", style='List Bullet')
        
        # Save document directly to generated_docs_sow folder
        project_name = context.get('project_info', 'SOW')
        if isinstance(project_name, str) and len(project_name) > 50:
            project_name = project_name[:50]
        
        filename = self._default_filename(project_name)
        out = self.out_dir / filename
        doc.save(str(out))
        print(f"✅ Enhanced SOW document generated: {out}")
        print(f"📄 Filename: {filename}")
        return out